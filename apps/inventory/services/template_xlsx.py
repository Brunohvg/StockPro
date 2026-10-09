"""
Gera o modelo de importação em Excel (.xlsx), com instruções e exemplos.

Abas:
- "Instruções": passo a passo e explicação de cada coluna.
- "Produtos" (ou "Movimentação"): onde o usuário preenche. Só o cabeçalho.
- "Exemplo": a mesma estrutura já preenchida (não é importada).
- "Listas": categorias, marcas, locais e fornecedores já cadastrados na empresa.
"""
import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill('solid', fgColor='4F46E5')
REQUIRED_FILL = PatternFill('solid', fgColor='B91C1C')
OPTIONAL_LOT_FILL = PatternFill('solid', fgColor='047857')
HEADER_FONT = Font(bold=True, color='FFFFFF')
TITLE_FONT = Font(bold=True, size=16, color='1E1B4B')
SUB_FONT = Font(bold=True, size=12, color='4F46E5')
THIN = Side(style='thin', color='CBD5E1')

# (coluna, obrigatória, largura, formato, explicação, exemplo)
CATALOG_SPEC = [
    ('nome', True, 34, None, "Nome do produto. Obrigatório.", "Chocolate ao Leite 90g"),
    ('sku', False, 16, '@', "Seu código interno. É a chave: se já existir, o produto é ATUALIZADO. "
     "Se ficar vazio, o sistema gera um código e procura pelo nome.", "CHOC-090"),
    ('categoria', False, 18, None, "Nome da categoria. Se não existir, é criada.", "Doces"),
    ('marca', False, 16, None, "Nome da marca. Se não existir, é criada.", "Cacau Bom"),
    ('unidade', False, 10, None, "UN, CX, KG, G, L, ML, M, PC, PAR, KIT. Padrão: UN.", "UN"),
    ('codigo_barras', False, 18, '@', "EAN/GTIN. A coluna já está como texto para não virar 7,89E+12.", "7891234567895"),
    ('custo', False, 12, '#,##0.00', "Custo unitário em R$ (aceita 12,50 ou 12.50).", 4.2),
    ('preco_venda', False, 12, '#,##0.00', "Preço de venda em R$. Usado no relatório de margem.", 7.9),
    ('estoque', False, 10, '#,##0.###', "Saldo atual (contagem). Substitui o saldo do sistema. "
     "Deixe vazio para não mexer no estoque.", 40),
    ('estoque_minimo', False, 14, '#,##0.###', "Abaixo disso o produto aparece como estoque baixo.", 10),
    ('validade', False, 13, 'DD/MM/YYYY', "Data de validade (dd/mm/aaaa). Preenchida, o produto passa a "
     "controlar validade. Mesmo produto com 2 validades: repita a linha com o mesmo SKU, "
     "uma por validade, cada uma com sua quantidade em 'estoque'.", date.today() + timedelta(days=120)),
    ('lote', False, 12, '@', "Número do lote (opcional).", "L2310A"),
    ('fabricacao', False, 13, 'DD/MM/YYYY', "Data de fabricação (opcional).", date.today() - timedelta(days=30)),
    ('fornecedor', False, 24, None, "Nome do fornecedor. Para criar um novo, informe também o CNPJ.", "Distribuidora Doce Ltda"),
    ('cnpj', False, 20, '@', "CNPJ do fornecedor (com ou sem pontuação).", ""),
    ('local', False, 16, None, "Local de estoque padrão (loja, depósito...). Se não existir, é criado.", "Loja"),
    ('sku_pai', False, 14, '@', "Só para produtos com variação (cor, tamanho): o mesmo sku_pai agrupa as "
     "variações num produto só.", ""),
    ('atributos', False, 22, None, "Atributos da variação, no formato Cor:Azul; Tamanho:M", ""),
    ('descricao_detalhada', False, 30, None, "Descrição longa do produto (opcional).", ""),
]

STOCK_SPEC = [
    ('sku', True, 16, '@', "SKU ou código de barras de um produto JÁ cadastrado.", "CHOC-090"),
    ('quantidade', True, 12, '#,##0.###', "Positivo = ENTRADA (soma). Negativo = SAÍDA (subtrai). "
     "Ex.: 24 ou -3.", 24),
    ('nome', False, 30, None, "Só para você se orientar; não altera o cadastro.", "Chocolate ao Leite 90g"),
    ('custo', False, 12, '#,##0.00', "Custo unitário da entrada (atualiza o custo médio).", 4.35),
    ('validade', False, 13, 'DD/MM/YYYY', "Validade do lote que está entrando (dd/mm/aaaa).", date.today() + timedelta(days=150)),
    ('lote', False, 12, '@', "Número do lote (opcional).", "L2311B"),
    ('fabricacao', False, 13, 'DD/MM/YYYY', "Data de fabricação (opcional).", ""),
]

EXTRA_CATALOG_EXAMPLES = [
    # mesmo SKU, segunda validade
    {'nome': 'Chocolate ao Leite 90g', 'sku': 'CHOC-090', 'categoria': 'Doces', 'estoque': 15,
     'validade': date.today() + timedelta(days=45), 'lote': 'L2309Z'},
    # produto sem validade
    {'nome': 'Fita de Cetim 10mm Vermelha', 'sku': 'FITA-10-VM', 'categoria': 'Aviamentos', 'marca': 'Progresso',
     'unidade': 'M', 'codigo_barras': '7898765432109', 'custo': 0.35, 'preco_venda': 0.9, 'estoque': 500,
     'estoque_minimo': 100, 'local': 'Depósito'},
    # produto com variação
    {'nome': 'Linha Amigurumi - Azul', 'sku': 'AMI-AZ', 'categoria': 'Linhas', 'marca': 'Círculo', 'unidade': 'UN',
     'custo': 14.9, 'preco_venda': 24.9, 'estoque': 12, 'sku_pai': 'AMI', 'atributos': 'Cor:Azul'},
    {'nome': 'Linha Amigurumi - Rosa', 'sku': 'AMI-RS', 'categoria': 'Linhas', 'marca': 'Círculo', 'unidade': 'UN',
     'custo': 14.9, 'preco_venda': 24.9, 'estoque': 8, 'sku_pai': 'AMI', 'atributos': 'Cor:Rosa'},
]

EXTRA_STOCK_EXAMPLES = [
    {'sku': 'FITA-10-VM', 'quantidade': -20, 'nome': 'Fita de Cetim 10mm Vermelha'},
]


def _write_header(ws, spec):
    for col, (name, required, width, fmt, help_text, _) in enumerate(spec, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = HEADER_FONT
        cell.fill = REQUIRED_FILL if required else (
            OPTIONAL_LOT_FILL if name in ('validade', 'lote', 'fabricacao') else HEADER_FILL)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.comment = Comment(help_text, 'StockPro', width=320, height=140)
        ws.column_dimensions[get_column_letter(col)].width = width
        if fmt:
            for row in range(2, 1002):
                ws.cell(row=row, column=col).number_format = fmt
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 22


def _write_rows(ws, spec, rows, start=2):
    names = [c[0] for c in spec]
    for r, data in enumerate(rows, start=start):
        for c, name in enumerate(names, start=1):
            value = data.get(name, '')
            if value != '':
                ws.cell(row=r, column=c, value=value)


def _add_validations(ws, spec):
    names = [c[0] for c in spec]
    if 'unidade' in names:
        col = get_column_letter(names.index('unidade') + 1)
        dv = DataValidation(type='list', formula1='"UN,CX,KG,G,L,ML,M,PC,PAR,KIT,FD,DZ"', allow_blank=True,
                            showErrorMessage=False)
        dv.add(f'{col}2:{col}1001')
        ws.add_data_validation(dv)
    for name in ('validade', 'fabricacao'):
        if name in names:
            col = get_column_letter(names.index(name) + 1)
            dv = DataValidation(type='date', operator='greaterThan', formula1='DATE(2000,1,1)', allow_blank=True,
                                error='Use uma data no formato dd/mm/aaaa.', errorTitle='Data inválida',
                                showErrorMessage=True)
            dv.add(f'{col}2:{col}1001')
            ws.add_data_validation(dv)
    for name in ('estoque', 'estoque_minimo', 'custo', 'preco_venda'):
        if name in names:
            col = get_column_letter(names.index(name) + 1)
            dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True,
                                error='Informe um número maior ou igual a zero.', showErrorMessage=True)
            dv.add(f'{col}2:{col}1001')
            ws.add_data_validation(dv)


def _instructions(ws, spec, kind):
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 100
    is_catalog = kind == 'CATALOG_DIRECT'
    sheet = 'Produtos' if is_catalog else 'Movimentação'
    ws['A1'] = 'Modelo de importação StockPro: ' + ('Catálogo' if is_catalog else 'Movimentação de estoque')
    ws['A1'].font = TITLE_FONT
    steps = (
        [
            f"1. Preencha a aba \"{sheet}\" (uma linha por produto). A aba \"Exemplo\" mostra como fica; ela não é importada.",
            "2. Só a coluna vermelha é obrigatória. As verdes são de validade/lote. Passe o mouse no cabeçalho para ver a ajuda.",
            "3. Salve como .xlsx (recomendado) ou CSV. Pode ser CSV com vírgula ou ponto e vírgula, com acentos.",
            "4. No StockPro: Importações > Nova importação > Carga de Catálogo > envie o arquivo.",
            "5. Confira a PRÉVIA (quantos produtos serão criados/atualizados, categorias novas e erros por linha). "
            "Nada é gravado até você clicar em Confirmar.",
        ] if is_catalog else [
            f"1. Preencha a aba \"{sheet}\": uma linha por produto que entrou ou saiu.",
            "2. Quantidade positiva soma (entrada); negativa subtrai (saída). Produtos com validade saem primeiro do lote "
            "que vence antes (FEFO).",
            "3. O produto precisa existir no catálogo (use o SKU ou o código de barras).",
            "4. No StockPro: Importações > Nova importação > Movimentação > envie o arquivo e confira a prévia.",
        ]
    )
    row = 3
    ws.cell(row=row, column=1, value='Como usar').font = SUB_FONT
    for text in steps:
        row += 1
        ws.cell(row=row, column=1, value=text)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
    row += 2
    ws.cell(row=row, column=1, value='Colunas').font = SUB_FONT
    row += 1
    for c, title in enumerate(('Coluna', 'Obrigatória', 'O que colocar'), start=1):
        cell = ws.cell(row=row, column=c, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for name, required, _w, _f, help_text, _e in spec:
        row += 1
        ws.cell(row=row, column=1, value=name).font = Font(bold=True)
        ws.cell(row=row, column=2, value='Sim' if required else 'Não')
        cell = ws.cell(row=row, column=3, value=help_text)
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        for c in range(1, 4):
            ws.cell(row=row, column=c).border = Border(bottom=THIN)
    if is_catalog:
        row += 2
        ws.cell(row=row, column=1, value='Dicas').font = SUB_FONT
        for text in (
            "• Atualizar preços ou categorias em massa: exporte/preencha só nome + sku + as colunas que quer mudar. "
            "Colunas vazias não apagam nada.",
            "• Validade: o mesmo produto com duas datas = duas linhas com o mesmo SKU, cada uma com sua quantidade. "
            "O sistema soma as linhas e guarda cada lote separado.",
            "• Os nomes de categoria/marca/local não diferenciam maiúsculas: 'doces' e 'Doces' são a mesma.",
            "• A aba \"Listas\" mostra o que já está cadastrado na sua empresa, para você usar os mesmos nomes.",
            "• Limite: 20.000 linhas e 10 MB por arquivo.",
        ):
            row += 1
            ws.cell(row=row, column=1, value=text)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)


def _lists(ws, tenant):
    from apps.inventory.models import Location
    from apps.partners.models import Supplier
    from apps.products.models import Brand, Category

    columns = [
        ('Categorias', Category.objects.filter(tenant=tenant).order_by('name').values_list('name', flat=True)),
        ('Marcas', Brand.objects.filter(tenant=tenant).order_by('name').values_list('name', flat=True)),
        ('Locais', Location.objects.filter(tenant=tenant, is_active=True).order_by('name').values_list('name', flat=True)),
        ('Fornecedores', Supplier.objects.filter(tenant=tenant).order_by('trade_name').values_list('trade_name', flat=True)),
    ]
    for c, (title, values) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(c)].width = 28
        for r, value in enumerate(values[:2000], start=2):
            ws.cell(row=r, column=c, value=value)
    ws.freeze_panes = 'A2'


def build_import_template(tenant, kind='CATALOG_DIRECT'):
    is_catalog = kind != 'STOCK_SYNC'
    spec = CATALOG_SPEC if is_catalog else STOCK_SPEC
    kind = 'CATALOG_DIRECT' if is_catalog else 'STOCK_SYNC'

    wb = Workbook()
    ws_help = wb.active
    ws_help.title = 'Instruções'
    _instructions(ws_help, spec, kind)

    ws_data = wb.create_sheet('Produtos' if is_catalog else 'Movimentação')
    _write_header(ws_data, spec)
    _add_validations(ws_data, spec)

    ws_example = wb.create_sheet('Exemplo')
    _write_header(ws_example, spec)
    first = {c[0]: c[5] for c in spec}
    _write_rows(ws_example, spec, [first] + (EXTRA_CATALOG_EXAMPLES if is_catalog else EXTRA_STOCK_EXAMPLES))

    if is_catalog and tenant is not None:
        _lists(wb.create_sheet('Listas'), tenant)

    wb.active = 1  # abre direto na aba de preenchimento
    buf = io.BytesIO()
    wb.save(buf)
    filename = 'modelo_catalogo_stockpro.xlsx' if is_catalog else 'modelo_movimentacao_stockpro.xlsx'
    return buf.getvalue(), filename
