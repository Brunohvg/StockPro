"""
Leitura de planilhas de importação (.xlsx e CSV) de forma robusta.

- .xlsx/.xlsm via openpyxl; CSV com ',' ';' tab ou '|', em UTF-8 ou Windows-1252.
- Tudo lido como texto: EAN/SKU nunca viram '789.0' nem notação científica.
- Cabeçalhos normalizados: minúsculas, sem acento, espaços viram '_'
  ("Código de Barras" -> "codigo_de_barras").
"""
import csv
import io
import os
import re
import unicodedata

import pandas as pd

MAX_ROWS = 20000


class SpreadsheetError(ValueError):
    pass


def normalize_header(name):
    text = unicodedata.normalize('NFKD', str(name)).encode('ascii', 'ignore').decode()
    text = re.sub(r'[^a-z0-9]+', '_', text.strip().lower()).strip('_')
    return text


def _read_csv_bytes(raw):
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise SpreadsheetError("Não foi possível ler o arquivo. Salve como CSV UTF-8 ou .xlsx.")

    sample = '\n'.join(text.splitlines()[:20])
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=',;\t|').delimiter
    except csv.Error:
        delimiter = ';' if sample.count(';') > sample.count(',') else ','
    df = pd.read_csv(io.StringIO(text), sep=delimiter, dtype=str, keep_default_na=False)
    # Undo the apostrophe escape used by StockPro CSV exports. Preserve the
    # literal text in the database; only exported CSV cells receive the prefix.
    for col in df.columns:
        df[col] = df[col].map(lambda v: v[1:] if isinstance(v, str) and v.startswith("'")
                             and (v[1:].lstrip().startswith(('=', '+', '-', '@'))
                                  or v[1:].startswith(('\t', '\r', '\n'))) else v)
    return df


def _cell_to_text(value):
    if value is None:
        return ''
    if hasattr(value, 'strftime'):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # EAN digitado como número no Excel
    return str(value).strip()


def _read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    # Usa a aba "Produtos" se existir; senão a primeira que não seja de instruções
    names = wb.sheetnames
    sheet_name = next((n for n in names if normalize_header(n) in ('produtos', 'products', 'estoque', 'movimentacao')), None)
    if sheet_name is None:
        sheet_name = next((n for n in names if not any(k in normalize_header(n) for k in ('instru', 'exemplo', 'listas'))), names[0])
    ws = wb[sheet_name]
    rows = ws.iter_rows(values_only=True)
    header = None
    data = []
    for row in rows:
        cells = [_cell_to_text(v) for v in row]
        if header is None:
            if any(cells):
                header = cells
            continue
        if any(cells):
            data.append(cells[:len(header)] + [''] * (len(header) - len(cells)))
        if len(data) > MAX_ROWS:
            break
    wb.close()
    if header is None:
        raise SpreadsheetError("A planilha está vazia.")
    return pd.DataFrame(data, columns=header, dtype=str)


def read_spreadsheet(path):
    ext = os.path.splitext(str(path))[1].lower()
    try:
        if ext in ('.xlsx', '.xlsm'):
            df = _read_xlsx(path)
        elif ext == '.xls':
            raise SpreadsheetError("Formato .xls antigo não é suportado. Salve como .xlsx.")
        else:
            with open(path, 'rb') as fh:
                df = _read_csv_bytes(fh.read())
    except SpreadsheetError:
        raise
    except Exception as exc:
        raise SpreadsheetError(f"Não foi possível ler a planilha: {exc}") from exc

    df.columns = [normalize_header(c) for c in df.columns]
    df = df.loc[:, [c for c in df.columns if c]]
    df = df.fillna('')
    df = df[(df != '').any(axis=1)].reset_index(drop=True)
    if len(df) > MAX_ROWS:
        raise SpreadsheetError(f"A planilha tem mais de {MAX_ROWS} linhas. Divida em arquivos menores.")
    return df
