"""Patch 3: importação por planilha (xlsx/CSV BR, prévia, modelo) e validade por lote (FEFO)."""
import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import TenantMembership
from apps.core.models import SystemSetting
from apps.core.services import StockService, parse_date_br
from apps.inventory.models import ImportBatch, MovementLot, StockLot
from apps.inventory.services.template_xlsx import build_import_template
from apps.inventory.tasks import process_csv_catalog_direct, process_csv_stock_adjustment
from apps.products.models import Category, Product, ProductVariant
from tests.factories import TenantFactory, UserFactory

PWD = 'Senha!Forte123'
TODAY = date.today()


def _member(role='OWNER'):
    tenant = TenantFactory()
    tenant.plan.max_products = 1000
    tenant.plan.save()
    u = UserFactory(password=PWD)
    TenantMembership.objects.create(user=u, tenant=tenant, role=role)
    return u, tenant


def _product(tenant, sku='LEITE', tracks=True):
    return Product.objects.create(tenant=tenant, name=f'Produto {sku}', sku=sku, tracks_expiry=tracks)


def _stock(p):
    return p.variants.first().current_stock


@pytest.fixture
def media(tmp_path, settings):
    settings.MEDIA_ROOT = str(tmp_path)
    return tmp_path


# ---------------------------------------------------------------- lotes / FEFO

@pytest.mark.django_db
class TestLotes:
    def test_saida_consumo_fefo(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 10, product=p, expiry_date=TODAY + timedelta(days=60), lot_number='B')
        StockService.create_movement(t, u, 'IN', 5, product=p, expiry_date=TODAY + timedelta(days=10), lot_number='A')
        mov = StockService.create_movement(t, u, 'OUT', 7, product=p)
        lots = {l.lot_number: l.quantity for l in StockLot.objects.filter(variant__product=p)}
        assert lots == {'A': 0, 'B': 8}
        assert _stock(p) == 8
        assert sorted((a.lot.lot_number, a.quantity) for a in mov.lot_allocations.all()) == [('A', 5), ('B', 2)]


    def test_lote_vencido_nao_pode_ser_consumido(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 5, product=p,
                                     expiry_date=TODAY - timedelta(days=1), lot_number='VENCIDO')
        with pytest.raises(ValueError, match='vencido'):
            StockService.create_movement(t, u, 'OUT', 1, product=p)
        lot = StockLot.objects.get(lot_number='VENCIDO')
        with pytest.raises(ValueError, match='vencido'):
            StockService.create_movement(t, u, 'OUT', 1, product=p, lot_id=lot.pk)
        assert _stock(p) == 5
        lot.refresh_from_db()
        assert lot.quantity == 5

    def test_fefo_ignora_vencido_e_preserva_o_seu_saldo(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 5, product=p,
                                     expiry_date=TODAY - timedelta(days=1), lot_number='VENCIDO')
        StockService.create_movement(t, u, 'IN', 3, product=p,
                                     expiry_date=TODAY + timedelta(days=3), lot_number='VALIDO')
        StockService.create_movement(t, u, 'OUT', 2, product=p)
        lots = {l.lot_number: l.quantity for l in StockLot.objects.filter(variant__product=p)}
        assert lots == {'VENCIDO': 5, 'VALIDO': 1}
        with pytest.raises(ValueError, match='vencidos'):
            StockService.create_movement(t, u, 'OUT', 2, product=p)
        assert _stock(p) == 6

    def test_saida_de_lote_especifico(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 5, product=p, expiry_date=TODAY + timedelta(days=10), lot_number='A')
        StockService.create_movement(t, u, 'IN', 5, product=p, expiry_date=TODAY + timedelta(days=90), lot_number='B')
        lot_b = StockLot.objects.get(lot_number='B')
        StockService.create_movement(t, u, 'OUT', 3, product=p, lot_id=lot_b.pk)
        lot_b.refresh_from_db()
        assert lot_b.quantity == 2
        with pytest.raises(ValueError):
            StockService.create_movement(t, u, 'OUT', 3, product=p, lot_id=lot_b.pk)

    def test_entrada_com_validade_liga_controle(self):
        u, t = _member()
        p = _product(t, tracks=False)
        StockService.create_movement(t, u, 'IN', 4, product=p, expiry_date='31/12/2030')
        p.refresh_from_db()
        assert p.tracks_expiry
        assert StockLot.objects.get(variant__product=p).expiry_date == date(2030, 12, 31)

    def test_saldo_sem_lote_sai_depois_dos_lotes(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 10, product=p)  # sem validade
        StockService.create_movement(t, u, 'IN', 4, product=p, expiry_date=TODAY + timedelta(days=5))
        StockService.create_movement(t, u, 'OUT', 6, product=p)
        assert StockLot.objects.get(variant__product=p).quantity == 0
        assert _stock(p) == 8

    def test_ajuste_com_lote_vira_contagem(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 10, product=p, expiry_date=TODAY + timedelta(days=5), lot_number='A')
        StockService.create_movement(t, u, 'ADJ', 7, product=p, expiry_date=TODAY + timedelta(days=50), lot_number='C')
        lots = {l.lot_number: l.quantity for l in StockLot.objects.all()}
        assert lots == {'A': 0, 'C': 7} and _stock(p) == 7

    def test_ajuste_sem_lote_reduz_os_que_vencem_antes(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 5, product=p, expiry_date=TODAY + timedelta(days=5), lot_number='A')
        StockService.create_movement(t, u, 'IN', 5, product=p, expiry_date=TODAY + timedelta(days=50), lot_number='B')
        StockService.create_movement(t, u, 'ADJ', 6, product=p)
        lots = {l.lot_number: l.quantity for l in StockLot.objects.all()}
        assert lots == {'A': 1, 'B': 5}

    def test_produto_sem_controle_nao_cria_lote(self):
        u, t = _member()
        p = _product(t, tracks=False)
        StockService.create_movement(t, u, 'IN', 5, product=p)
        assert not StockLot.objects.exists() and not MovementLot.objects.exists()

    def test_datas(self):
        assert parse_date_br('05/03/2027') == date(2027, 3, 5)
        assert parse_date_br('2027-03-05 00:00:00') == date(2027, 3, 5)
        assert parse_date_br('') is None
        with pytest.raises(ValueError):
            parse_date_br('31/31/2027')
        u, t = _member()
        with pytest.raises(ValueError):
            StockService.create_movement(t, u, 'IN', 1, product=_product(t),
                                         expiry_date='01/01/2027', manufacture_date='01/02/2027')


# ---------------------------------------------------------------- importação

def _batch(tenant, user, content, name, kind='CATALOG_DIRECT'):
    return ImportBatch.objects.create(tenant=tenant, user=user, type=kind,
                                      file=SimpleUploadedFile(name, content))


def _xlsx_from_template(tenant, rows, kind='CATALOG_DIRECT'):
    content, _ = build_import_template(tenant, kind)
    wb = load_workbook(io.BytesIO(content))
    ws = wb['Produtos' if kind == 'CATALOG_DIRECT' else 'Movimentação']
    header = [c.value for c in ws[1]]
    for r, data in enumerate(rows, start=2):
        for c, name in enumerate(header, start=1):
            if name in data:
                ws.cell(row=r, column=c, value=data[name])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


@pytest.mark.django_db
class TestImportacao:
    def test_modelo_xlsx_tem_instrucoes_exemplo_e_listas(self):
        _, t = _member()
        Category.objects.create(tenant=t, name='Doces')
        content, filename = build_import_template(t)
        wb = load_workbook(io.BytesIO(content))
        assert filename.endswith('.xlsx')
        assert wb.sheetnames == ['Instruções', 'Produtos', 'Exemplo', 'Listas']
        assert wb['Produtos']['A1'].value == 'nome' and wb['Produtos']['A2'].value is None
        assert wb['Listas']['A2'].value == 'Doces'

    def test_modelo_nao_importa_aba_de_exemplo(self, media):
        u, t = _member()
        content, _ = build_import_template(t)
        log = process_csv_catalog_direct(_batch(t, u, content, 'modelo.xlsx'))
        assert Product.objects.filter(tenant=t).count() == 0, log

    def test_xlsx_completo_com_validade(self, media):
        u, t = _member()
        content = _xlsx_from_template(t, [
            {'nome': 'Leite Condensado', 'sku': 'LC-395', 'categoria': 'Mercearia', 'marca': 'Moça',
             'codigo_barras': '7891000100103', 'custo': 5.5, 'preco_venda': 8.99, 'estoque': 10,
             'estoque_minimo': 4, 'validade': TODAY + timedelta(days=20), 'lote': 'A1'},
            {'nome': 'Leite Condensado', 'sku': 'LC-395', 'estoque': 6,
             'validade': TODAY + timedelta(days=90), 'lote': 'A2'},
            {'nome': 'Fita Cetim', 'sku': '000123', 'categoria': 'aviamentos', 'estoque': 50},
        ])
        log = process_csv_catalog_direct(_batch(t, u, content, 'carga.xlsx'))
        assert '0 erros' in log, log
        p = Product.objects.get(tenant=t, sku='LC-395')
        v = p.variants.first()
        assert p.category.name == 'Mercearia' and p.brand.name == 'Moça'
        assert p.tracks_expiry and p.sale_price == Decimal('8.99')
        assert v.barcode == '7891000100103' and v.minimum_stock == 4 and v.current_stock == 16
        assert {l.lot_number: l.quantity for l in v.lots.all()} == {'A1': 10, 'A2': 6}
        assert ProductVariant.objects.get(tenant=t, sku='000123').current_stock == 50  # zero à esquerda mantido

    def test_csv_do_excel_ponto_e_virgula_e_acentos(self, media):
        u, t = _member()
        csv_bytes = ("Nome;Código de Barras;Categoria;Estoque;Preço Venda\n"
                     "Pão de Mel;7890000000017;Doces;5;3,50\n"
                     "Paçoca;;Doces;2;1,00\n"
                     "Bala;789;Doces;;\n").encode('cp1252')
        log = process_csv_catalog_direct(_batch(t, u, csv_bytes, 'excel.csv'))
        assert '3 criados' in log and '0 erros' in log, log
        assert ProductVariant.objects.get(tenant=t, product__name='Bala').barcode == '789'
        assert Product.objects.get(tenant=t, name='Pão de Mel').sale_price == Decimal('3.50')
        assert Category.objects.filter(tenant=t).count() == 1

    def test_erros_por_linha_nao_param_a_carga(self, media):
        u, t = _member()
        csv_bytes = b"nome,sku,estoque,validade\nOk,OK1,1,\n,SEMNOME,1,\nData ruim,DR1,1,31/31/2026\n"
        log = process_csv_catalog_direct(_batch(t, u, csv_bytes, 'x.csv'))
        assert '1 criados' in log and '2 erros' in log
        assert 'Linha 3' in log and 'Linha 4' in log

    def test_movimentacao_por_planilha_usa_fefo(self, media):
        u, t = _member()
        p = _product(t, 'IOG')
        StockService.create_movement(t, u, 'IN', 5, product=p, expiry_date=TODAY + timedelta(days=3), lot_number='V')
        content = _xlsx_from_template(t, [
            {'sku': 'IOG', 'quantidade': 10, 'validade': TODAY + timedelta(days=40), 'lote': 'N'},
            {'sku': 'IOG', 'quantidade': -6},
        ], kind='STOCK_SYNC')
        log = process_csv_stock_adjustment(_batch(t, u, content, 'mov.xlsx', 'STOCK_SYNC'))
        assert '0 erros' in log, log
        assert {l.lot_number: l.quantity for l in StockLot.objects.all()} == {'V': 0, 'N': 9}

    def test_previa_nao_grava_e_confirmar_grava(self, media, monkeypatch):
        u, t = _member()
        c = Client()
        c.force_login(u)
        from apps.inventory import tasks
        monkeypatch.setattr(tasks.process_import_task, 'delay', lambda bid: tasks.process_import_task(bid))
        f = SimpleUploadedFile('p.csv', b"nome,sku,categoria,estoque\nCaneta,CAN,Papelaria,3\n")
        r = c.post('/inventory/imports/new/', {'type': 'CATALOG_DIRECT', 'file': f})
        batch = ImportBatch.objects.get(tenant=t)
        assert r.status_code == 302 and batch.status == 'PENDING_REVIEW'
        assert 'PRÉVIA' in batch.log and '1 criados' in batch.log
        assert not Product.objects.filter(tenant=t).exists()
        assert not Category.objects.filter(tenant=t).exists()
        c.post(f'/inventory/imports/{batch.pk}/confirm/')
        batch.refresh_from_db()
        assert batch.status == 'COMPLETED', batch.log
        assert Product.objects.get(tenant=t, sku='CAN').category.name == 'Papelaria'

    def test_arquivo_ilegivel_vira_erro(self, media):
        u, t = _member()
        c = Client()
        c.force_login(u)
        f = SimpleUploadedFile('p.csv', b"\xff\xfe\x00\x00lixo")
        c.post('/inventory/imports/new/', {'type': 'CATALOG_DIRECT', 'file': f})
        assert ImportBatch.objects.get(tenant=t).status == 'ERROR'

    def test_extensao_invalida_recusada(self, media):
        u, t = _member()
        c = Client()
        c.force_login(u)
        c.post('/inventory/imports/new/', {'type': 'CATALOG_DIRECT', 'file': SimpleUploadedFile('a.pdf', b'%PDF')})
        assert not ImportBatch.objects.filter(tenant=t).exists()

    def test_download_do_modelo(self):
        u, _ = _member()
        c = Client()
        c.force_login(u)
        r = c.get('/inventory/imports/template.xlsx?type=STOCK_SYNC')
        assert r.status_code == 200
        assert load_workbook(io.BytesIO(r.content)).sheetnames == ['Instruções', 'Movimentação', 'Exemplo']


# ---------------------------------------------------------------- telas, API e alertas

@pytest.mark.django_db
class TestValidadeTelasEApi:
    def test_painel_mostra_vencidos_e_vencendo(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 2, product=p, expiry_date=TODAY - timedelta(days=1))
        StockService.create_movement(t, u, 'IN', 3, product=p, expiry_date=TODAY + timedelta(days=7))
        c = Client()
        c.force_login(u)
        html = c.get('/app/').content.decode()
        assert 'Vencidos (1)' in html and 'Vencem em até 30 dias (1)' in html

    def test_detalhe_do_produto_lista_lotes(self):
        u, t = _member()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 2, product=p, lot_number='LT-77', expiry_date=TODAY + timedelta(days=7))
        c = Client()
        c.force_login(u)
        html = c.get(f'/products/{p.pk}/').content.decode()
        assert 'Lotes e validade' in html and 'LT-77' in html

    def test_tela_de_movimentacao_registra_validade(self):
        u, t = _member()
        p = _product(t, tracks=False)
        c = Client()
        c.force_login(u)
        c.post('/inventory/movements/add/', {'product_identifier': p.sku, 'type': 'IN', 'quantity': '3',
                                              'expiry_date': '2031-01-15', 'lot_number': 'X9'})
        lot = StockLot.objects.get(variant__product=p)
        assert lot.expiry_date == date(2031, 1, 15) and lot.lot_number == 'X9' and lot.quantity == 3

    def test_mobile_registra_validade_na_entrada(self):
        u, t = _member('OPERATOR')
        p = _product(t)
        c = Client()
        c.force_login(u)
        c.post('/mobile/move/', {'type': 'IN', 'sku': p.sku, 'quantity': '4', 'expiry_date': '2031-02-01'})
        assert StockLot.objects.get(variant__product=p).quantity == 4

    def test_api_entrada_com_validade_e_busca_mostra_proximo_vencimento(self):
        u, t = _member()
        p = _product(t)
        api = APIClient()
        tok = api.post('/api/v1/auth/token/', {'username': u.username, 'password': PWD}, format='json').data['access']
        api.credentials(HTTP_AUTHORIZATION=f'Bearer {tok}')
        exp = (TODAY + timedelta(days=12)).isoformat()
        r = api.post('/api/v1/inventory/entry/', {'items': [
            {'sku': p.sku, 'quantity': 5, 'expiry_date': exp, 'lot_number': 'API1'}]}, format='json')
        assert r.status_code == 200, r.data
        res = api.get(f'/api/v1/products/search/?q={p.sku}').data['results'][0]
        assert res['next_expiry'] == exp and res['next_expiry_days'] == 12

    def test_alerta_diario_por_email(self):
        from apps.inventory.services.expiry import send_expiry_alerts
        u, t = _member()
        cfg = SystemSetting.get_settings(t)
        cfg.alert_email = 'estoque@empresa.com'
        cfg.save()
        p = _product(t)
        StockService.create_movement(t, u, 'IN', 2, product=p, expiry_date=TODAY + timedelta(days=3))
        assert send_expiry_alerts() == 1
        assert mail.outbox[0].to == ['estoque@empresa.com'] and p.sku in mail.outbox[0].body

    def test_configuracoes_salvam_dias_de_aviso(self):
        u, t = _member()
        c = Client()
        c.force_login(u)
        c.post('/settings/', {'company_name': 'X', 'low_stock_alert_threshold': 5, 'expiry_alert_days': 45})
        assert SystemSetting.get_settings(t).expiry_alert_days == 45
