"""Patch 10: nome das variações, exportação no padrão do import, etiqueta Bibelô, números da Visão Geral e do BI."""
import csv
import io
import json
from decimal import Decimal

import pytest
from django.core.files.base import ContentFile
from django.core.management import call_command
from django.test import Client
from django.urls import reverse
from openpyxl import load_workbook

from apps.accounts.models import TenantMembership
from apps.core.services import StockService
from apps.inventory.models import ImportBatch, Location, StockMovement
from apps.inventory.tasks import run_import
from apps.partners.models import Supplier
from apps.products.models import (
    AttributeType, Brand, Category, Product, ProductType, ProductVariant, VariantAttributeValue,
)
from apps.reports.exports import ProductExporter
from tests.factories import TenantFactory, UserFactory

pytestmark = pytest.mark.django_db
PWD = 'Senha!Forte123'


def _member(role='OWNER', tenant=None):
    tenant = tenant or TenantFactory()
    u = UserFactory(password=PWD)
    TenantMembership.objects.create(user=u, tenant=tenant, role=role)
    return u, tenant


def _client(u):
    c = Client()
    c.force_login(u)
    return c


def _variable(t, name='Tinta Chalk Paint - Acrilex 100 ml', sku='VAR-TIN-0011'):
    return Product.objects.create(tenant=t, name=name, sku=sku, product_type=ProductType.VARIABLE)


def _variant(p, sku, attrs, name=''):
    v = ProductVariant.objects.create(tenant=p.tenant, product=p, sku=sku, name=name)
    for key, val in attrs.items():
        at, _ = AttributeType.objects.get_or_create(tenant=p.tenant, name=key)
        VariantAttributeValue.objects.create(variant=v, attribute_type=at, value=val)
    return v


# ---------------------------------------------------------------- nome das variações

class TestNomeVariacao:
    def test_nova_variacao_pela_tela_recebe_nome_e_sku_dos_atributos(self):
        u, t = _member()
        p = _variable(t)
        cor = AttributeType.objects.create(tenant=t, name='Cor')
        r = _client(u).post(reverse('products:variant_create', args=[p.pk]), {
            'sku': '', 'name': f'{p.name} - ', 'current_stock': 0, 'minimum_stock': 0,
            'is_active': 'on', f'attr_{cor.id}': 'Azul',
        })
        assert r.status_code == 302
        v = p.variants.get()
        assert v.name == 'Tinta Chalk Paint - Acrilex 100 ml - Azul'
        assert v.sku == 'VAR-TIN-0011-AZ'

    def test_formulario_nao_vem_mais_com_nome_provisorio(self):
        u, t = _member()
        p = _variable(t)
        r = _client(u).get(reverse('products:variant_create', args=[p.pk]))
        assert f'value="{p.name} - "' not in r.content.decode()

    def test_nome_digitado_pelo_usuario_fica(self):
        u, t = _member()
        p = _variable(t)
        cor = AttributeType.objects.create(tenant=t, name='Cor')
        _client(u).post(reverse('products:variant_create', args=[p.pk]), {
            'sku': 'X-1', 'name': 'Azul Royal especial', 'current_stock': 0, 'minimum_stock': 0,
            f'attr_{cor.id}': 'Azul',
        })
        v = p.variants.get()
        assert v.name == 'Azul Royal especial' and v.sku == 'X-1'

    def test_editar_atributo_atualiza_nome_automatico(self):
        u, t = _member()
        p = _variable(t)
        v = _variant(p, 'V-1', {'Cor': 'Azul'})
        v.refresh_auto_name()
        cor = AttributeType.objects.get(tenant=t, name='Cor')
        _client(u).post(reverse('products:variant_edit', args=[v.pk]), {
            'sku': 'V-1', 'name': v.name, 'current_stock': 0, 'minimum_stock': 0,
            f'attr_{cor.id}': 'Verde',
        })
        v.refresh_from_db()
        assert v.name.endswith(' - Verde')

    def test_comando_corrige_nomes_antigos(self):
        _, t = _member()
        p = _variable(t)
        a = _variant(p, 'VAR-TIN-0011-11', {'Cor': 'Azul'}, name='Tinta Chalk Paint - Acrilex 100 ml -')
        b = _variant(p, 'VAR-TIN-0011-12', {'Cor': 'Rosa'}, name='Rosinha da loja')
        out = io.StringIO()
        call_command('fix_variant_names', stdout=out)
        a.refresh_from_db()
        assert a.name.endswith('-')  # sem --apply só mostra
        assert '1 variação' in out.getvalue()
        call_command('fix_variant_names', '--apply', stdout=io.StringIO())
        a.refresh_from_db(); b.refresh_from_db()
        assert a.name == 'Tinta Chalk Paint - Acrilex 100 ml - Azul'
        assert b.name == 'Rosinha da loja'


# ---------------------------------------------------------------- exportação

def _catalogo(t, u):
    cat = Category.objects.create(tenant=t, name='Tinta Pva')
    brand = Brand.objects.create(tenant=t, name='Acrilex')
    sup = Supplier.objects.create(tenant=t, company_name='Acrilex Tintas Ltda', trade_name='Acrilex',
                                  cnpj='11222333000181')
    loc = Location.objects.create(tenant=t, name='Loja', code='LOJA')
    p = _variable(t)
    p.category, p.brand, p.default_supplier, p.default_location = cat, brand, sup, loc
    p.description, p.sale_price = 'Tinta à base de água', Decimal('19.90')
    p.save()
    a = _variant(p, 'VAR-TIN-0011-AZ', {'Cor': 'Azul'})
    b = _variant(p, 'VAR-TIN-0011-RS', {'Cor': 'Rosa', 'Tamanho': 'G'})
    b.sale_price, b.barcode, b.minimum_stock = Decimal('21.50'), '7891234567895', Decimal('5')
    b.save()
    StockService.create_movement(t, u, 'IN', Decimal('40'), variant=b, unit_cost=Decimal('8.3333'))
    s = Product.objects.create(tenant=t, name='Caneta gel; azul "fina"', sku='CAN-1', category=cat,
                               sale_price=Decimal('4.50'), uom='UN')
    sv = s.variants.first()
    sv.barcode = '0789123456789'
    sv.minimum_stock = Decimal('2.5')
    sv.save()
    StockService.create_movement(t, u, 'IN', Decimal('1250.5'), variant=sv, unit_cost=Decimal('1.2'))
    return p, s


def _snapshot(t):
    out = []
    for v in ProductVariant.objects.filter(tenant=t).select_related('product').order_by('sku'):
        p = v.product
        out.append((v.sku, v.display_name, p.product_type, p.sku, p.name, v.current_stock, v.minimum_stock,
                    v.avg_unit_cost, v.sale_price if v.sale_price is not None else p.sale_price,
                    v.barcode or '', p.category_id, p.brand_id, p.default_supplier_id, p.default_location_id,
                    p.description, p.uom,
                    tuple(sorted((a.attribute_type.name, a.value) for a in v.attribute_values.all()))))
    return out


class TestExportacao:
    def test_csv_tem_exatamente_as_colunas_do_modelo(self):
        from apps.inventory.services.template_xlsx import CATALOG_SPEC
        u, t = _member()
        _catalogo(t, u)
        text = ProductExporter(t).export_csv()
        rows = list(csv.reader(io.StringIO(text.lstrip('﻿')), delimiter=';'))
        assert rows[0] == [c[0] for c in CATALOG_SPEC]
        by_sku = {r[1]: dict(zip(rows[0], r)) for r in rows[1:]}
        assert set(by_sku) == {'VAR-TIN-0011-AZ', 'VAR-TIN-0011-RS', 'CAN-1'}  # pai não vira linha
        rosa = by_sku['VAR-TIN-0011-RS']
        assert rosa['nome'] == 'Tinta Chalk Paint - Acrilex 100 ml - Rosa / G'
        assert rosa['sku_pai'] == 'VAR-TIN-0011'
        assert rosa['atributos'] == 'Cor:Rosa; Tamanho:G'
        assert rosa['estoque'] == '40' and rosa['custo'] == '8,3333' and rosa['preco_venda'] == '21,5'
        assert rosa['fornecedor'] == 'Acrilex' and rosa['cnpj'] == '11222333000181'
        assert by_sku['VAR-TIN-0011-AZ']['preco_venda'] == '19,9'  # herda o preço do pai
        assert by_sku['CAN-1']['sku_pai'] == '' and by_sku['CAN-1']['estoque'] == '1250,5'

    @pytest.mark.parametrize('fmt', ['csv', 'xlsx'])
    def test_exportar_e_importar_de_volta_nao_muda_nada(self, fmt):
        u, t = _member()
        _catalogo(t, u)
        before = _snapshot(t)
        movements = StockMovement.objects.filter(tenant=t).count()
        exp = ProductExporter(t)
        content = exp.export_csv().encode('utf-8') if fmt == 'csv' else exp.export_excel()
        batch = ImportBatch.objects.create(tenant=t, user=u, type='CATALOG_DIRECT')
        batch.file.save(f'export.{fmt}', ContentFile(content))
        result = run_import(batch)
        assert 'Linha' not in result or 'erro' not in result.lower(), result
        assert _snapshot(t) == before
        assert StockMovement.objects.filter(tenant=t).count() == movements
        assert Product.objects.filter(tenant=t).count() == 2

    def test_importar_exportacao_em_outra_empresa_recria_o_catalogo(self):
        u, t = _member()
        _catalogo(t, u)
        u2, t2 = _member()
        batch = ImportBatch.objects.create(tenant=t2, user=u2, type='CATALOG_DIRECT')
        batch.file.save('export.xlsx', ContentFile(ProductExporter(t).export_excel()))
        run_import(batch)
        p = Product.objects.get(tenant=t2, sku='VAR-TIN-0011')
        assert p.product_type == ProductType.VARIABLE and p.name == 'Tinta Chalk Paint - Acrilex 100 ml'
        rosa = ProductVariant.objects.get(tenant=t2, sku='VAR-TIN-0011-RS')
        assert rosa.current_stock == 40 and rosa.display_name.endswith('Rosa / G')

    def test_xlsx_tem_aba_produtos_e_codigo_como_texto(self):
        u, t = _member()
        _catalogo(t, u)
        wb = load_workbook(io.BytesIO(ProductExporter(t).export_excel()))
        ws = wb['Produtos']
        header = [c.value for c in ws[1]]
        row = {header[i]: c.value for i, c in enumerate(ws[2])}
        assert row['sku'] == 'CAN-1' and row['codigo_barras'] == '0789123456789'

    def test_json_traz_preco_fornecedor_e_custo(self):
        u, t = _member()
        p, _ = _catalogo(t, u)
        Product.objects.filter(pk=p.pk).update(avg_unit_cost=Decimal('7.10'))
        data = {i['sku']: i for i in json.loads(ProductExporter(t).export_json())}
        var = data['VAR-TIN-0011']
        assert var['sale_price'] == 19.9
        assert var['supplier'] == {'name': 'Acrilex', 'cnpj': '11222333000181'}
        variants = {v['sku']: v for v in var['variants']}
        assert variants['VAR-TIN-0011-RS']['cost'] == 8.3333
        assert variants['VAR-TIN-0011-AZ']['cost'] == 7.1  # sem custo na variação: custo do produto
        assert variants['VAR-TIN-0011-AZ']['name'].endswith(' - Azul')
        assert data['CAN-1']['sale_price'] == 4.5


# ---------------------------------------------------------------- etiqueta "código + nome + EAN"

from apps.labels import zpl as zpl_mod  # noqa: E402
from apps.labels.barcodes import ean13_modules  # noqa: E402
from apps.labels.layout import Barcode, LabelItem, Text, build_label  # noqa: E402
from apps.labels.models import LabelSettings, VariantLabel  # noqa: E402
from apps.labels.preview import render_svg  # noqa: E402


def _cfg_bibelo(dpi=203):
    return LabelSettings(width_mm=33, height_mm=22, columns=3, column_gap_mm=Decimal('2'), dpi=dpi,
                         layout='code_name', code_label='CÓDIGO:')


APLIQUE = LabelItem('Aplique flor prensada', '', None, 'APL-FL', '7890000139557', '', code='139557',
                    title='Aplique flor prensada - PCT 50 UN')


class TestEtiquetaCodigoNome:
    @pytest.mark.parametrize('dpi', [203, 300])
    def test_desenho_igual_a_foto(self, dpi):
        label = build_label(APLIQUE, _cfg_bibelo(dpi))
        texts = [e.text for e in label.elements if isinstance(e, Text)]
        assert texts[:3] == ['CÓDIGO: 139557', 'APLIQUE FLOR', 'PRENSADA - PCT 50 UN']
        assert texts[3:] == ['7', '890000', '139557']  # 1º dígito fora, dois grupos de 6
        assert not label.warnings
        for el in label.elements:
            w = el.width if isinstance(el, Text) else len(el.modules) * el.module
            h = el.size if isinstance(el, Text) else el.height + el.guard_ext
            assert 0 <= el.x and el.x + w <= label.width and el.y + h <= label.height

    def test_barras_de_guarda_mais_longas_e_codigo_intacto(self):
        label = build_label(APLIQUE, _cfg_bibelo())
        bar = next(e for e in label.elements if isinstance(e, Barcode))
        assert bar.modules == ean13_modules('7890000139557')
        bars = bar.bars()
        assert sum(1 for _, _, h in bars if h > bar.height) == 6  # 2 + 2 + 2 guardas
        # as barras desenhadas reconstroem os 95 módulos
        mods = ['0'] * 95
        for x, w, _ in bars:
            for k in range((x - bar.x) // bar.module, (x - bar.x + w) // bar.module):
                mods[k] = '1'
        assert ''.join(mods) == bar.modules

    def test_zpl_e_previa_tem_os_mesmos_elementos(self):
        cfg = _cfg_bibelo()
        labels = [build_label(APLIQUE, cfg)] * 3
        out = zpl_mod.render_row(labels, cfg)
        assert out.count('^FDCÓDIGO: 139557^FS') == 3
        assert out.count('^FD890000^FS') == 3
        svg = render_svg(labels, cfg)
        assert svg.count('CÓDIGO: 139557') == 3

    def test_sem_texto_antes_do_codigo_e_nome_comprido(self):
        cfg = _cfg_bibelo()
        cfg.code_label = ''
        item = LabelItem('x', '', None, 'SKU-1', '', '', title='Fita de cetim 10mm vermelha rolo 50 metros progresso')
        label = build_label(item, cfg)
        texts = [e.text for e in label.elements if isinstance(e, Text)]
        assert texts[0] == 'SKU-1'
        assert texts[2].endswith('...') and any('Edite o nome' in w for w in label.warnings)

    def test_estilo_completo_continua_igual(self):
        cfg = LabelSettings(width_mm=50, height_mm=30, dpi=203)
        label = build_label(APLIQUE, cfg)
        texts = [e.text for e in label.elements if isinstance(e, Text)]
        assert 'Aplique flor prensada' in texts and 'CÓDIGO: 139557' not in texts


class TestTextoDaEtiqueta:
    def _setup(self):
        u, t = _member('OPERATOR')
        p = Product.objects.create(tenant=t, name='Aplique flor prensada - PCT 50 UN', sku='APL-FL')
        v = p.variants.first()
        v.barcode = '7890000139557'
        v.save()
        cfg = LabelSettings.for_tenant(t)
        cfg.layout, cfg.width_mm, cfg.height_mm, cfg.columns = 'code_name', 33, 22, 3
        cfg.save()
        return u, t, v

    def test_salva_nome_e_codigo_e_usa_na_impressao(self):
        u, t, v = self._setup()
        c = _client(u)
        r = c.post(reverse('labels:save_text', args=[v.pk]), {'name': 'Aplique flor', 'code': '139557'})
        item = r.json()['item']
        assert (item['label_name'], item['label_code']) == ('Aplique flor', '139557')
        assert VariantLabel.objects.get(variant=v).updated_by == u
        out = c.post(reverse('labels:generate'), {f'q_{v.pk}': '3', 'formato': 'zpl'}).content.decode()
        assert '^FDCÓDIGO: 139557^FS' in out and '^FDAPLIQUE FLOR^FS' in out
        # a busca já devolve o texto gravado
        found = c.get(reverse('labels:search'), {'q': 'aplique'}).json()['results'][0]
        assert found['label_code'] == '139557'

    def test_padrao_e_voltar_ao_cadastro(self):
        u, t, v = self._setup()
        c = _client(u)
        out = c.post(reverse('labels:generate'), {f'q_{v.pk}': '1', 'formato': 'zpl'}).content.decode()
        assert '^FDCÓDIGO: APL-FL^FS' in out  # sem edição: SKU
        c.post(reverse('labels:save_text', args=[v.pk]), {'name': 'X', 'code': '1'})
        c.post(reverse('labels:save_text', args=[v.pk]), {'name': '', 'code': ''})
        assert not VariantLabel.objects.filter(variant=v).exists()

    def test_outra_empresa_nao_edita(self):
        _, _, v = self._setup()
        intruso, _ = _member()
        r = _client(intruso).post(reverse('labels:save_text', args=[v.pk]), {'name': 'Hack', 'code': ''})
        assert r.status_code == 404
        assert not VariantLabel.objects.exists()

    def test_tamanho_maximo(self):
        u, _, v = self._setup()
        r = _client(u).post(reverse('labels:save_text', args=[v.pk]), {'name': 'x' * 81, 'code': ''})
        assert r.status_code == 400

    def test_modelo_salva_estilo(self):
        u, t = _member('ADMIN')
        r = _client(u).post('/etiquetas/modelo/', {
            'layout': 'code_name', 'code_label': 'CÓD.', 'width_mm': 33, 'height_mm': 22, 'columns': 3,
            'column_gap_mm': 2, 'dpi': 203, 'darkness': 0, 'print_speed': 0, 'offset_x_mm': 0,
            'offset_y_mm': 0, 'show_barcode': 'on', 'code_source': 'auto'})
        assert r.status_code == 302
        cfg = LabelSettings.objects.get(tenant=t)
        assert (cfg.layout, cfg.code_label, cfg.columns) == ('code_name', 'CÓD.', 3)


# ---------------------------------------------------------------- números (metrics)

from datetime import timedelta  # noqa: E402

from django.utils import timezone  # noqa: E402

from apps.inventory.services.stock_alerts import weekly_summary  # noqa: E402
from apps.products.services import ProductArchiveService  # noqa: E402
from apps.reports import metrics  # noqa: E402
from apps.reports.templatetags.metrics_tags import brl, qtd  # noqa: E402


def _loja():
    """Caneta (simples, com categoria), Linha (variável, sem categoria) e vendas em dias diferentes."""
    u, t = _member()
    cat = Category.objects.create(tenant=t, name='Papelaria')
    caneta = Product.objects.create(tenant=t, name='Caneta', sku='CAN', category=cat, sale_price=Decimal('10'))
    StockService.create_movement(t, u, 'IN', 50, product=caneta, unit_cost=Decimal('4'))
    linha = _variable(t, name='Linha', sku='LIN')
    azul = _variant(linha, 'LIN-AZ', {'Cor': 'Azul'})
    azul.sale_price = Decimal('20')
    azul.minimum_stock = Decimal('5')
    azul.save()
    StockService.create_movement(t, u, 'IN', 10, variant=azul, unit_cost=Decimal('12'))
    StockService.create_movement(t, u, 'OUT', 6, product=caneta)   # venda: 60 de receita, 24 de custo
    StockService.create_movement(t, u, 'OUT', 6, variant=azul)     # venda: 120 de receita, 72 de custo
    old = StockService.create_movement(t, u, 'OUT', 2, product=caneta)
    StockMovement.objects.filter(pk=old.pk).update(created_at=timezone.now() - timedelta(days=40))
    return u, t, caneta, azul


class TestNumeros:
    def test_mesmo_numero_na_visao_geral_no_bi_e_no_email(self):
        u, t, caneta, azul = _loja()
        overview = metrics.overview(t)
        bi = metrics.intelligence(t, metrics.period('7'))
        email = weekly_summary(t)
        # 42 canetas × 4 + 4 linhas × 12 = 216
        assert overview['stock']['value'] == bi['stock']['value'] == email['stock_value'] == Decimal('216')
        assert bi['sales']['units'] == email['out_qty'] == Decimal('12')
        html_dash = _client(u).get('/app/').content.decode()
        html_bi = _client(u).get('/app/analytics/?periodo=7').content.decode()
        assert 'R$ 216' in html_dash and 'R$ 216' in html_bi

    def test_campos_antigos_do_produto_nao_entram(self):
        _, t, caneta, _ = _loja()
        Product.objects.filter(pk=caneta.pk).update(current_stock=999, avg_unit_cost=100)
        assert metrics.stock_snapshot(t)['value'] == Decimal('216')

    def test_arquivamento_nao_e_venda_e_arquivado_sai_do_estoque(self):
        u, t, caneta, _ = _loja()
        ProductArchiveService.remove_product(caneta, u)
        start, end = metrics.period('7').bounds()
        assert metrics.sales_summary(t, start, end)['units'] == Decimal('12')
        assert metrics.stock_snapshot(t)['value'] == Decimal('48')

    def test_vendas_cmv_margem_e_comparacao(self):
        _, t, _, _ = _loja()
        bi = metrics.intelligence(t, metrics.period('30'))
        s = bi['sales']
        assert (s['units'], s['revenue'], s['cogs'], s['margin']) == (12, Decimal('180'), Decimal('96'), Decimal('84'))
        assert round(s['margin_pct']) == 47
        assert bi['prev']['units'] == 2  # a venda de 40 dias atrás cai no período anterior
        assert round(bi['changes']['units']) == 500
        assert bi['coverage_days'] == int(Decimal('216') / (Decimal('96') / 30))

    def test_grafico_tem_todos_os_dias(self):
        _, t, _, _ = _loja()
        per = metrics.period('30')
        series = metrics.daily_series(t, per)
        assert len(series) == 30
        assert series[-1]['out'] == 12 and series[-1]['in'] == 60
        assert sum(d['out'] for d in series[:-1]) == 0

    def test_categoria_inclui_sem_categoria(self):
        _, t, _, _ = _loja()
        cats = {c['name']: c['value'] for c in metrics.category_values(t)}
        assert cats == {'Papelaria': Decimal('168'), 'Sem categoria': Decimal('48')}

    def test_abc_pelo_faturamento(self):
        _, t, _, _ = _loja()
        abc = metrics.intelligence(t, metrics.period('30'))['abc']
        # Linha (R$ 120) e Caneta (R$ 60): a Caneta é quem cruza os 80%, então também é A
        assert [r['variant'].sku for r in abc['A']['items']] == ['LIN-AZ', 'CAN']
        assert abc['B']['count'] == 0 and abc['total'] == Decimal('180')
        assert abc['A']['pct'] == 100.0

    def test_parados_e_quadro_de_atencao(self):
        u, t, caneta, azul = _loja()
        parado = Product.objects.create(tenant=t, name='Cola', sku='COL')
        StockService.create_movement(t, u, 'IN', 3, product=parado, unit_cost=Decimal('5'))
        stall = metrics.stalled(t)
        assert [v.sku for v in stall['items']] == ['COL'] and stall['value'] == Decimal('15')
        StockService.create_movement(t, u, 'OUT', 4, variant=azul)  # chega a 0 e vendeu: zerado que vende
        keys = {a['key']: a['count'] for a in metrics.attention(t)}
        assert keys['zeroed'] == 1
        assert keys['no_price'] == 1 and keys['no_barcode'] == 3

    def test_empresa_nova_ve_primeiros_passos(self):
        u, t = _member()
        html = _client(u).get('/app/').content.decode()
        assert 'Primeiros passos' in html and '0 de 5' in html
        Product.objects.create(tenant=t, name='X', sku='X')
        assert metrics.onboarding(t)['done'] == 1

    def test_bi_por_periodo_e_sem_ao_vivo(self):
        u, _, _, _ = _loja()
        c = _client(u)
        for key in ('7', '30', '90', 'mes'):
            r = c.get(f'/app/analytics/?periodo={key}')
            assert r.status_code == 200
        html = c.get('/app/analytics/?periodo=x').content.decode()  # inválido cai em 30 dias
        assert 'aria-current="page">30 dias' in html and 'Ao Vivo' not in html

    def test_quantidade_sem_casas_desnecessarias(self):
        assert qtd(Decimal('240.0000')) == '240'
        assert qtd(Decimal('2.5000')) == '2,5'
        assert qtd(Decimal('1250')) == '1.250'
        assert brl(Decimal('1234.5')) == 'R$ 1.234,50' and brl(Decimal('1234.5'), 0) == 'R$ 1.234'

    def test_periodo_mes_compara_com_mesmo_trecho_do_mes_anterior(self):
        from datetime import date
        per = metrics.period('mes', today=date(2026, 10, 10))
        assert (per.start, per.end) == (date(2026, 10, 1), date(2026, 10, 10))
        assert (per.prev_start, per.prev_end) == (date(2026, 9, 1), date(2026, 9, 10))


class TestFiltrosDoCatalogo:
    def test_sem_preco_sem_codigo_e_baixo_pelas_variacoes(self):
        u, t, caneta, azul = _loja()
        c = _client(u)
        sem_preco = Product.objects.create(tenant=t, name='Cola', sku='COL')
        html = c.get('/products/?falta=preco').content.decode()
        assert 'Cola' in html and 'Caneta' not in html
        html = c.get('/products/?stock=low').content.decode()
        assert 'Linha' in html and 'Caneta' not in html  # Linha: 4 com mínimo 5
        assert sem_preco
